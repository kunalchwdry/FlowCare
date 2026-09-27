"""Approximate raster preview of a deck (DejaVu stands in for the real fonts).

Recurses into groups, draws solid fills, pastes pictures, and lays out text so
overlaps and overflow are visible. Geometry is approximate — use it to catch
collisions, not to judge typography.
"""
import io
import sys
from math import ceil

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu, Length
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN

SCALE = 70
REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
BLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
_f = {}


def font(size, bold):
    k = (round(size, 1), bold)
    if k not in _f:
        _f[k] = ImageFont.truetype(BLD if bold else REG, max(5, int(size * SCALE / 72)))
    return _f[k]


def px(v):
    return Emu(int(v)).inches * SCALE


def rgb(c, d=(0, 0, 0)):
    try:
        return tuple(bytes.fromhex(str(c)))
    except Exception:
        return d


def wrap(d, text, f, width):
    out = []
    for para in text.split("\n"):
        line = ""
        for w in para.split(" "):
            t = (line + " " + w).strip()
            if d.textlength(t, font=f) <= width or not line:
                line = t
            else:
                out.append(line)
                line = w
        out.append(line)
    return out


def draw_shape(d, img, sh, ox=0, oy=0, sx=1.0, sy=1.0):
    if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
        gx, gy = px(sh.left), px(sh.top)
        gw, gh = px(sh.width), px(sh.height)
        ns = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
        xfrm = sh._element.find('.//' + ns + 'xfrm')
        if xfrm is None:
            return
        chOff, chExt = xfrm.find(ns + 'chOff'), xfrm.find(ns + 'chExt')
        cx, cy = int(chOff.get('x')), int(chOff.get('y'))
        cw, ch = int(chExt.get('cx')), int(chExt.get('cy'))
        nsx = gw / max(px(cw), 1e-6)
        nsy = gh / max(px(ch), 1e-6)
        for c in sh.shapes:
            draw_shape(d, img, c,
                       ox + gx - px(cx) * nsx, oy + gy - px(cy) * nsy,
                       nsx, nsy)
        return

    x = ox + px(sh.left) * sx
    y = oy + px(sh.top) * sy
    w = px(sh.width) * sx
    h = px(sh.height) * sy

    if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
        try:
            im = Image.open(io.BytesIO(sh.image.blob)).convert("RGBA")
            im = im.resize((max(1, int(w)), max(1, int(h))))
            img.paste(im, (int(x), int(y)), im)
        except Exception:
            pass
        return

    try:
        if sh.fill.type == 1:
            d.rectangle([x, y, x + w, y + h], fill=rgb(sh.fill.fore_color.rgb, (210, 210, 210)))
    except Exception:
        pass
    try:
        if sh.line.fill.type == 1:
            d.rectangle([x, y, x + w, y + h], outline=rgb(sh.line.color.rgb, (160, 160, 160)))
    except Exception:
        pass

    if not sh.has_text_frame:
        return
    ty = y
    for para in sh.text_frame.paragraphs:
        runs = [r for r in para.runs if r.text]
        if not runs:
            ty += 6
            continue
        sz = max((r.font.size.pt if r.font.size else 14) for r in runs) * min(sx, sy)
        bold = any(r.font.bold for r in runs)
        col = rgb(next((r.font.color.rgb for r in runs
                        if r.font.color and r.font.color.type is not None), None))
        f = font(sz, bold)
        ls = para.line_spacing
        lh = (ls.pt if isinstance(ls, Length) else sz * (ls or 1.2)) * SCALE / 72
        for line in wrap(d, "".join(r.text for r in runs), f, max(10, w)):
            lx = x
            if para.alignment == PP_ALIGN.CENTER:
                lx = x + (w - d.textlength(line, font=f)) / 2
            elif para.alignment == PP_ALIGN.RIGHT:
                lx = x + w - d.textlength(line, font=f)
            d.text((lx, ty), line, font=f, fill=col)
            ty += lh
        sa = para.space_after
        ty += (sa.pt if isinstance(sa, Length) else 0) * SCALE / 72


src = sys.argv[1] if len(sys.argv) > 1 else "FlowCare-CuriousParc-2026.pptx"
out = sys.argv[2] if len(sys.argv) > 2 else "preview"
prs = Presentation(src)
W, H = int(px(prs.slide_width)), int(px(prs.slide_height))
for i, slide in enumerate(prs.slides, 1):
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    for sh in slide.shapes:
        draw_shape(d, img, sh)
    img.save(f"{out}/slide-{i:02d}.png")
print("rendered", len(prs.slides._sldIdLst), "slides ->", out)
